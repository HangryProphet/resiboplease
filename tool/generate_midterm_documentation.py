"""Generate the repository-backed Resibo, Please midterm documentation DOCX.

The script deliberately uses only facts verified from the current v1.0 source,
schema, tests, and tracked documentation. It also generates compact diagram PNGs
that are embedded in the final Word document.
"""

from __future__ import annotations

import math
import textwrap
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = ROOT / "docs" / "midterm" / "assets"
OUTPUT = ROOT / "docs" / "RESIBO_PLEASE_MIDTERM_DOCUMENTATION.docx"

NAVY = "12283B"
DEEP_NAVY = "081725"
CREAM = "FFF4D6"
PAPER = "F4E3B4"
GOLD = "E6B84A"
TEAL = "2B8C87"
RED = "B94755"
BLUE = "3975A6"
GREEN = "4B8456"
GRAY = "E9ECEF"
MID_GRAY = "65707A"
INK = "1B2733"
WHITE = "FFFFFF"

FONT_REGULAR = Path("C:/Windows/Fonts/arial.ttf")
FONT_BOLD = Path("C:/Windows/Fonts/arialbd.ttf")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT_REGULAR), size)


def wrapped_lines(text: str, width: int, size: int, bold: bool = False) -> list[str]:
    probe = Image.new("RGB", (10, 10))
    draw = ImageDraw.Draw(probe)
    used_font = font(size, bold)
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=used_font)[2] <= width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    lines: list[str] | None = None,
    *,
    title_size: int = 42,
    body_size: int = 28,
    title_color: str = INK,
    body_color: str = INK,
) -> None:
    x1, y1, x2, y2 = box
    max_width = x2 - x1 - 50
    title_lines = wrapped_lines(title, max_width, title_size, True)
    body_lines: list[str] = []
    for line in lines or []:
        body_lines.extend(wrapped_lines(line, max_width, body_size, False))
    title_height = len(title_lines) * (title_size + 8)
    body_height = len(body_lines) * (body_size + 7)
    gap = 18 if body_lines else 0
    total_height = title_height + gap + body_height
    y = y1 + ((y2 - y1) - total_height) / 2
    for value in title_lines:
        bbox = draw.textbbox((0, 0), value, font=font(title_size, True))
        draw.text(
            ((x1 + x2 - (bbox[2] - bbox[0])) / 2, y),
            value,
            font=font(title_size, True),
            fill=f"#{title_color}",
        )
        y += title_size + 8
    y += gap
    for value in body_lines:
        bbox = draw.textbbox((0, 0), value, font=font(body_size))
        draw.text(
            ((x1 + x2 - (bbox[2] - bbox[0])) / 2, y),
            value,
            font=font(body_size),
            fill=f"#{body_color}",
        )
        y += body_size + 7


def rounded_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    lines: list[str] | None = None,
    *,
    fill: str = CREAM,
    outline: str = NAVY,
    title_color: str = INK,
    body_color: str = INK,
    radius: int = 26,
    width: int = 5,
    title_size: int = 42,
    body_size: int = 28,
) -> None:
    draw.rounded_rectangle(
        box,
        radius=radius,
        fill=f"#{fill}",
        outline=f"#{outline}",
        width=width,
    )
    draw_centered_text(
        draw,
        box,
        title,
        lines,
        title_size=title_size,
        body_size=body_size,
        title_color=title_color,
        body_color=body_color,
    )


def arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    *,
    color: str = NAVY,
    width: int = 8,
) -> None:
    draw.line([start, end], fill=f"#{color}", width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 22
    left = (
        end[0] - length * math.cos(angle - math.pi / 6),
        end[1] - length * math.sin(angle - math.pi / 6),
    )
    right = (
        end[0] - length * math.cos(angle + math.pi / 6),
        end[1] - length * math.sin(angle + math.pi / 6),
    )
    draw.polygon([end, left, right], fill=f"#{color}")


def diagram_canvas(width: int, height: int, title: str, subtitle: str) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((20, 20, width - 20, height - 20), radius=34, outline=f"#{NAVY}", width=5)
    draw.rectangle((20, 20, width - 20, 180), fill=f"#{NAVY}")
    draw.text((70, 58), title, font=font(54, True), fill=f"#{CREAM}")
    draw.text((72, 124), subtitle, font=font(27), fill="#D7E1E8")
    return image, draw


def generate_architecture_diagram(path: Path) -> None:
    image, draw = diagram_canvas(
        1800,
        1780,
        "SYSTEM ARCHITECTURE",
        "Local-first deterministic game with an optional community service",
    )
    rounded_box(draw, (650, 225, 1150, 350), "PLAYER", ["Android / Flutter Web"], fill=GOLD)
    arrow(draw, (900, 350), (900, 410))
    rounded_box(
        draw,
        (350, 410, 1450, 575),
        "FLUTTER PRESENTATION & NAVIGATION",
        ["GoRouter • custom responsive screens • localization • audio • accessibility"],
        fill=CREAM,
    )
    arrow(draw, (900, 575), (900, 635))
    rounded_box(
        draw,
        (250, 635, 1550, 875),
        "GAME CONTROLLER & DOMAIN LAYER",
        [
            "Procedural City Generator • Candidate Generator • Investigation / Evidence",
            "Election • Term Engine • Living City Engine • structured models",
        ],
        fill="DCE9EC",
        outline=TEAL,
    )
    arrow(draw, (900, 875), (900, 935))
    rounded_box(
        draw,
        (210, 935, 1590, 1115),
        "PRESENTATION CONSUMERS OF SHARED GAME TRUTH",
        ["City Brief • Candidate Hub • Chronicle • News • CityBook • TV • Map • Term Report"],
        fill="F7E9C7",
        outline=GOLD,
    )
    arrow(draw, (650, 1115), (500, 1200))
    arrow(draw, (1150, 1115), (1300, 1200))
    rounded_box(
        draw,
        (120, 1200, 830, 1465),
        "LOCAL PERSISTENCE",
        [
            "Sembast: 5 versioned city slots",
            "SharedPreferences: settings and gallery caches",
            "Generated world, investigation, vote, and 365-day term state",
        ],
        fill="E8F1EB",
        outline=GREEN,
    )
    rounded_box(
        draw,
        (970, 1200, 1680, 1465),
        "OPTIONAL ONLINE COMMUNITY",
        [
            "HTTP API service → Supabase PostgREST",
            "Anonymous Auth ownership",
            "player_profiles + sanitized city_snapshots",
        ],
        fill="E7EDF5",
        outline=BLUE,
    )
    rounded_box(
        draw,
        (275, 1530, 1525, 1695),
        "CORE RULE",
        ["Algorithm creates truth; presentation displays that truth."],
        fill=NAVY,
        outline=GOLD,
        title_color=GOLD,
        body_color=WHITE,
        title_size=38,
        body_size=34,
    )
    image.save(path, quality=95)


def generate_flow_diagram(path: Path) -> None:
    image, draw = diagram_canvas(
        1800,
        2820,
        "USER AND SYSTEM FLOW",
        "Current one-election, one-year gameplay loop",
    )
    center_x1, center_x2 = 540, 1260
    y = 220
    steps = [
        ("LAUNCH APP", ["Immersive Android game shell"], GOLD, NAVY, 125),
        ("MAIN MENU", ["Continue City or Start New City"], CREAM, NAVY, 145),
        ("CITY SETUP & PROCEDURAL GENERATION", ["World seed • 4–6 issues • 5 candidates"], "DCE9EC", TEAL, 165),
        ("PRE-ELECTION CITY BRIEF", ["City conditions • issue files • Chronicle"], CREAM, NAVY, 155),
        ("CANDIDATE INVESTIGATION", ["Bio • Promises • Debate • Records • Receipts"], "F7E9C7", GOLD, 175),
        ("ELECTION", ["Candidate • priority issue • confidence • confirmation"], "F4DDE1", RED, 155),
        ("WINNER REVEAL", ["Elected mayor begins the term"], "F4DDE1", RED, 140),
        ("LIVING CITY SIMULATION", ["Interactive map • projects • events • shared media"], "DCE9EC", TEAL, 165),
        ("365-DAY TERM", ["Pause • Day • Week • Month • Year"], "E8F1EB", GREEN, 145),
        ("TERM REPORT", ["Indicators • gains • setbacks • history • context"], CREAM, NAVY, 155),
        ("OPTIONAL PUBLISH", ["Public or Private sanitized city snapshot"], "E7EDF5", BLUE, 145),
        ("VISIT CITIES", ["Explore • My Cities • player search • privacy/delete"], "E7EDF5", BLUE, 155),
    ]
    boxes: list[tuple[int, int, int, int]] = []
    for title, lines, fill_color, outline_color, height in steps:
        box = (center_x1, y, center_x2, y + height)
        rounded_box(
            draw,
            box,
            title,
            lines,
            fill=fill_color,
            outline=outline_color,
            title_size=36,
            body_size=25,
        )
        boxes.append(box)
        y += height + 55
    for first, second in zip(boxes, boxes[1:]):
        arrow(draw, ((first[0] + first[2]) // 2, first[3]), ((second[0] + second[2]) // 2, second[1]))

    rounded_box(
        draw,
        (70, 380, 430, 650),
        "MENU SUPPORT",
        ["How to Play", "Settings", "Visit Cities"],
        fill="EDF0F2",
        outline=MID_GRAY,
        title_size=32,
        body_size=25,
    )
    arrow(draw, (540, 445), (430, 500), color=MID_GRAY, width=6)
    rounded_box(
        draw,
        (1370, 735, 1730, 1100),
        "INVESTIGATION MEDIA",
        ["CityBook", "Chronicle", "News articles", "R-8 TV reports", "Evidence reader"],
        fill="EDF0F2",
        outline=MID_GRAY,
        title_size=32,
        body_size=24,
    )
    arrow(draw, (1260, 840), (1370, 840), color=MID_GRAY, width=6)
    rounded_box(
        draw,
        (70, 1730, 430, 2060),
        "SAVED PROGRESS",
        ["Return to menu", "Continue selected slot", "Resume pre-election, term, or report"],
        fill="EDF0F2",
        outline=MID_GRAY,
        title_size=32,
        body_size=23,
    )
    arrow(draw, (540, 1830), (430, 1880), color=MID_GRAY, width=6)
    draw.text(
        (420, 2740),
        "Online sharing is optional; the local city save remains the canonical game state.",
        font=font(27, True),
        fill=f"#{NAVY}",
    )
    image.save(path, quality=95)


def field_table_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    name: str,
    fields: list[str],
    *,
    color: str,
) -> None:
    x1, y1, x2, y2 = box
    draw.rounded_rectangle(box, radius=24, fill="#FFFFFF", outline=f"#{color}", width=6)
    draw.rounded_rectangle((x1, y1, x2, y1 + 100), radius=22, fill=f"#{color}")
    draw.rectangle((x1, y1 + 70, x2, y1 + 100), fill=f"#{color}")
    draw.text((x1 + 36, y1 + 26), name, font=font(39, True), fill="#FFFFFF")
    y = y1 + 125
    for item in fields:
        draw.text((x1 + 40, y), item, font=font(25), fill=f"#{INK}")
        y += 48


def generate_erd_diagram(path: Path) -> None:
    image, draw = diagram_canvas(
        1800,
        1650,
        "SUPABASE DATABASE ERD",
        "Current V2 online community schema",
    )
    field_table_box(
        draw,
        (90, 260, 620, 650),
        "auth.users",
        ["PK  id : uuid", "Supabase-managed identity", "Anonymous Auth session"],
        color=NAVY,
    )
    field_table_box(
        draw,
        (1040, 220, 1710, 720),
        "player_profiles",
        [
            "PK/FK  user_id : uuid",
            "display_name : text",
            "display_name_normalized : text",
            "created_at : timestamptz",
            "updated_at : timestamptz",
        ],
        color=TEAL,
    )
    arrow(draw, (620, 450), (1040, 450), color=TEAL, width=8)
    draw.text((710, 385), "1", font=font(32, True), fill=f"#{NAVY}")
    draw.text((945, 385), "0..1", font=font(32, True), fill=f"#{TEAL}")
    draw.text((700, 480), "profile identity", font=font(25), fill=f"#{MID_GRAY}")

    field_table_box(
        draw,
        (420, 855, 1380, 1480),
        "city_snapshots",
        [
            "PK  id : uuid        FK  owner_id : uuid",
            "is_public : boolean  •  city_name : text",
            "term_number : integer  •  starting_pressure : text",
            "top_issue : text  •  selected_candidate_name : text",
            "candidate_party : text  •  final_city_condition : text",
            "outcome_summary : text",
            "major_issues : jsonb[]  •  major_events : jsonb[]",
            "created_at : timestamptz",
        ],
        color=BLUE,
    )
    arrow(draw, (1375, 720), (1240, 855), color=BLUE, width=8)
    draw.text((1390, 742), "1", font=font(32, True), fill=f"#{TEAL}")
    draw.text((1260, 805), "0..*", font=font(32, True), fill=f"#{BLUE}")
    draw.text((1390, 795), "owns", font=font(25), fill=f"#{MID_GRAY}")
    draw.text(
        (330, 1540),
        "Local save slots are not database rows. Only sanitized completed-city summaries are uploaded.",
        font=font(28, True),
        fill=f"#{NAVY}",
    )
    image.save(path, quality=95)


def generate_component_diagram(path: Path) -> None:
    image, draw = diagram_canvas(
        1800,
        1850,
        "DATA AND COMPONENT ARCHITECTURE",
        "How one seed becomes consistent information and consequences",
    )
    rounded_box(draw, (650, 220, 1150, 345), "WORLD SEED", ["versioned and path-derived"], fill=GOLD)
    arrow(draw, (780, 345), (500, 430), color=TEAL)
    arrow(draw, (1020, 345), (1300, 430), color=TEAL)
    rounded_box(
        draw,
        (120, 430, 820, 655),
        "CITY ISSUE GENERATOR",
        ["Selects 4–6 of 12 issue families", "severity • urgency • trend • indicators"],
        fill="DCE9EC",
        outline=TEAL,
        title_size=36,
        body_size=26,
    )
    rounded_box(
        draw,
        (980, 430, 1680, 655),
        "CANDIDATE GENERATOR",
        ["Five unique candidates", "truth • history • promises • evidence"],
        fill="F7E9C7",
        outline=GOLD,
        title_size=36,
        body_size=26,
    )
    arrow(draw, (470, 655), (470, 750), color=TEAL)
    arrow(draw, (1330, 655), (1330, 750), color=GOLD)
    rounded_box(
        draw,
        (120, 750, 820, 975),
        "CITY PRESENTATION",
        ["City Brief • issue files", "Chronicle • community voices"],
        fill=CREAM,
        title_size=36,
        body_size=26,
    )
    rounded_box(
        draw,
        (980, 750, 1680, 975),
        "CANDIDATE PRESENTATION",
        ["Biography • Promises • Debate", "CityBook • News • Evidence"],
        fill=CREAM,
        title_size=36,
        body_size=26,
    )
    arrow(draw, (470, 975), (780, 1090), color=NAVY)
    arrow(draw, (1330, 975), (1020, 1090), color=NAVY)
    rounded_box(
        draw,
        (500, 1090, 1300, 1260),
        "PLAYER VOTE + ELECTED CANDIDATE",
        ["same saved city truth and candidate truth"],
        fill="F4DDE1",
        outline=RED,
        title_size=35,
        body_size=26,
    )
    arrow(draw, (900, 1260), (900, 1340), color=RED)
    rounded_box(
        draw,
        (420, 1340, 1380, 1515),
        "LIVING CITY ENGINE",
        ["365 deterministic days • projects • events • indicator changes"],
        fill="E8F1EB",
        outline=GREEN,
        title_size=38,
        body_size=27,
    )
    arrow(draw, (900, 1515), (900, 1590), color=GREEN)
    rounded_box(
        draw,
        (160, 1590, 1640, 1760),
        "ONE SAVED EVENT TRUTH → MANY PRESENTATIONS",
        ["Map • Event History • Chronicle • three News framings • TV • CityBook • Term Report"],
        fill=NAVY,
        outline=GOLD,
        title_color=GOLD,
        body_color=WHITE,
        title_size=36,
        body_size=26,
    )
    image.save(path, quality=95)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    # Ask Word-compatible editors to refresh the TOC and page-number fields on
    # open. The field also carries cached results after the final Word pass.
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings.element.append(update_fields)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color, before, after in [
        ("Title", 30, NAVY, 0, 12),
        ("Subtitle", 14, MID_GRAY, 0, 8),
        ("Heading 1", 20, NAVY, 12, 8),
        ("Heading 2", 14, TEAL, 9, 5),
        ("Heading 3", 11.5, RED, 7, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Figure Caption" not in styles:
        caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Figure Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("RESIBO, PLEASE  |  MIDTERM PROJECT DOCUMENTATION")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_run = footer.add_run("Page ")
    page_run.font.name = "Arial"
    page_run.font.size = Pt(8)
    add_field(footer, "PAGE", "1")

    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def add_paragraph(doc: Document, text: str, *, bold_lead: str | None = None, align=None):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        first.bold = True
        paragraph.add_run(text[len(bold_lead) :])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(doc: Document, items: list[str], *, level: int = 0) -> None:
    for item in items:
        style = "List Bullet" if level == 0 else "List Bullet 2"
        paragraph = doc.add_paragraph(item, style=style)
        paragraph.paragraph_format.space_after = Pt(3)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(item, style="List Number")
        paragraph.paragraph_format.space_after = Pt(4)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for index, label in enumerate(headers):
        cell = header_cells[index]
        cell.text = label
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9)
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string(WHITE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = value
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if row_index % 2 == 1:
                set_cell_shading(cell, "F4F6F7")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor.from_string(INK)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc: Document, title: str, text: str, *, color: str = TEAL) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EFF5F4")
    set_cell_margins(cell, top=130, bottom=130, start=160, end=160)
    first = cell.paragraphs[0]
    run = first.add_run(title.upper())
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(color)
    body = cell.add_paragraph(text)
    body.paragraph_format.space_after = Pt(0)
    body.runs[0].font.name = "Arial"
    body.runs[0].font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc: Document, image_path: Path, caption: str, width: float = 6.6) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    caption_paragraph = doc.add_paragraph(caption, style="Figure Caption")
    caption_paragraph.paragraph_format.keep_with_next = False


def add_section_heading(doc: Document, number: int, title: str) -> None:
    doc.add_heading(f"{number}. {title}", level=1)


def add_cover(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(28)
    paragraph.add_run().add_picture(
        str(ROOT / "assets" / "app_icon" / "resibo_launcher.png"),
        width=Inches(1.35),
    )
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(20)
    run = title.add_run("RESIBO, PLEASE")
    run.font.name = "Arial"
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(4)
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Midterm Project Documentation")
    run.font.name = "Arial"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(RED)

    formal = doc.add_paragraph()
    formal.alignment = WD_ALIGN_PARAGRAPH.CENTER
    formal.paragraph_format.left_indent = Inches(0.6)
    formal.paragraph_format.right_indent = Inches(0.6)
    formal_run = formal.add_run(
        "Resibo, Please: An AI-Assisted Mobile Civic Decision Simulation Game "
        "for Voter Literacy and SDG-Based Decision-Making"
    )
    formal_run.italic = True
    formal_run.font.name = "Arial"
    formal_run.font.size = Pt(12)
    formal_run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    doc.add_paragraph()
    identity = doc.add_table(rows=5, cols=2)
    identity.alignment = WD_TABLE_ALIGNMENT.CENTER
    identity.style = "Table Grid"
    details = [
        ("Student / Group", "[To be supplied before submission]"),
        ("Course / Section", "[To be supplied before submission]"),
        ("Application Version", "1.0.0+1"),
        ("Repository State", "Branch v1.0 • commit 561e7d6"),
        ("Document Date", "September 2, 2026"),
    ]
    for row, (label, value) in zip(identity.rows, details):
        set_cell_shading(row.cells[0], NAVY)
        set_cell_margins(row.cells[0], top=100, bottom=100)
        set_cell_margins(row.cells[1], top=100, bottom=100)
        row.cells[0].text = label
        row.cells[1].text = value
        for run in row.cells[0].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
            run.font.size = Pt(9)
        for run in row.cells[1].paragraphs[0].runs:
            run.font.size = Pt(9)

    notice = doc.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice.paragraph_format.space_before = Pt(22)
    notice_run = notice.add_run(
        "All politicians, cities, parties, media outlets, and events in this project are fictional."
    )
    notice_run.bold = True
    notice_run.font.size = Pt(9.5)
    notice_run.font.color.rgb = RGBColor.from_string(RED)
    notice2 = doc.add_paragraph()
    notice2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice2.add_run("The application does not provide real election guidance or candidate recommendations.").font.size = Pt(9)
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    doc.add_heading("Table of Contents", level=1)
    paragraph = doc.add_paragraph()
    add_field(
        paragraph,
        'TOC \\o "1-3" \\h \\z \\u',
        "The table of contents will update automatically when this document opens.",
    )
    doc.add_paragraph()
    add_callout(
        doc,
        "Document coverage",
        "This single file covers all six items under A. Project Documentation: title and description, objectives, scope and limitations, updated architecture/flowchart, database design, and the current progress report.",
        color=BLUE,
    )
    doc.add_page_break()


def build_document() -> Document:
    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = "Resibo, Please — Midterm Project Documentation"
    props.subject = "Project documentation for the Resibo, Please Flutter application"
    props.author = "Resibo, Please Project Team"
    props.keywords = "Flutter, civic simulation, voter literacy, Supabase, midterm"

    add_cover(doc)
    add_toc(doc)

    add_section_heading(doc, 1, "Project Title and Description")
    doc.add_heading("1.1 Project identity", level=2)
    add_table(
        doc,
        ["Item", "Verified project information"],
        [
            ["Project title", "Resibo, Please"],
            ["Formal title", "Resibo, Please: An AI-Assisted Mobile Civic Decision Simulation Game for Voter Literacy and SDG-Based Decision-Making"],
            ["Primary platform", "Android mobile; Flutter Web is also used for development and testing"],
            ["Technology", "Flutter/Dart with GoRouter, Sembast, SharedPreferences, HTTP, Supabase Anonymous Auth/PostgREST, Audioplayers, and a lightweight Flame menu layer"],
            ["Current delivery", "One complete fictional mayoral election followed by one deterministic 365-day term"],
        ],
        widths=[1.45, 5.05],
    )

    doc.add_heading("1.2 Project description", level=2)
    add_paragraph(
        doc,
        "Resibo, Please is a fictional civic and political decision simulation designed to let players practice evidence-based voting in a safe environment. A new local city is created from a seed and configuration, producing four to six active problems and five fictional mayoral candidates. The player studies the city, investigates the candidates, casts a ballot without being shown a recommended answer, and then observes how the elected mayor governs during a simulated 365-day year."
    )
    add_paragraph(
        doc,
        "Candidate checking is the central learning experience. Each generated candidate is represented through a Civic Archive biography, campaign promises, structured history, Campaign Trail material, CityBook posts, debate responses, comparative news coverage, public and campaign records, evidence relationships, unresolved questions, and a personal Receipts Board. These views are generated from one saved candidate-truth structure so that claims, records, and consequences remain connected."
    )
    add_paragraph(
        doc,
        "After voting, the Living City engine evaluates a complete deterministic year. City indicators, issue severity, promise-based projects, public trust, budget health, corruption pressure, world conditions, and causal event chains evolve from the inherited city and the elected candidate’s hidden abilities. The player can observe the city map, event history, Chronicle, News, R-8 television reports, CityBook responses, and the final Term Report."
    )
    add_paragraph(
        doc,
        "Local play remains the canonical experience. After a completed term, the player may optionally publish a deliberately sanitized Public or Private city snapshot to a Supabase community gallery. Other players can explore public snapshots, search display names, and manage their own uploaded snapshots; full saves and hidden simulation data are never part of the gallery payload."
    )

    doc.add_heading("1.3 Educational purpose and SDG alignment", level=2)
    add_table(
        doc,
        ["Alignment", "Project contribution"],
        [
            ["Primary: SDG 4 — Quality Education", "Supports civic education, media and information literacy, critical thinking, evidence evaluation, and reflective decision-making."],
            ["Secondary: SDG 11 — Sustainable Cities and Communities", "Uses local governance scenarios involving transport, services, resilience, housing-related infrastructure, health, education, water, employment, and public trust."],
            ["Related issue domains", "City issue records may also reference SDGs 1, 2, 3, 6, 8, and 13 when the generated problem is relevant."],
        ],
        widths=[2.1, 4.4],
    )
    add_paragraph(
        doc,
        "The intended primary users are senior high school students, college students, first-time or young voters, and civic education learners. Teachers and general mobile-game players are secondary audiences. The current game measures decisions inside a fictional simulation; it does not test or predict a real election."
    )
    add_callout(
        doc,
        "Political neutrality",
        "All candidates, parties, cities, records, media organizations, and events are fictional. The game never displays a best-candidate label, hidden quality score, match percentage, or voting recommendation and does not provide official election guidance.",
        color=RED,
    )

    add_section_heading(doc, 2, "Project Objectives")
    doc.add_heading("2.1 General objective", level=2)
    add_paragraph(
        doc,
        "To develop a mobile civic decision simulation that strengthens voter, media, and evidence literacy by allowing players to investigate fictional candidates, make an independent electoral decision, and observe the systemic consequences of that decision within a deterministic city simulation."
    )
    doc.add_heading("2.2 Specific objectives", level=2)
    add_numbered(
        doc,
        [
            "Generate varied but reproducible fictional cities containing four to six active problems selected from twelve civic issue families.",
            "Generate exactly five distinct fictional mayoral candidates per normal run and maintain consistent identities, histories, promises, evidence, media, and hidden tradeoffs through versioned seeded generation.",
            "Provide multiple evidence-facing investigation formats—including biographies, promises, debates, records, CityBook, Chronicle, News, television reports, and a Receipts Board—without exposing a pre-election recommendation.",
            "Allow the player to cast and confirm a ballot based on a selected candidate, priority issue, and self-reported confidence.",
            "Simulate one complete 365-day administration using bounded indicators, issue state, candidate abilities, projects, world conditions, and procedural events that remain stable after save restoration.",
            "Persist up to five local city runs, mutable investigation state, settings, election results, and Living City history using versioned storage and migrations.",
            "Provide English and Filipino interface support together with reduced-motion, high-contrast, audio, volume, and critical-event interruption settings.",
            "Implement an optional privacy-limited Supabase community feature with explicit REST GET, POST, PATCH, and DELETE operations for player profiles and sanitized city snapshots.",
            "Maintain responsive, game-styled interfaces that work on narrow Android screens and provide controlled scrolling instead of clipping on short displays.",
        ],
    )

    add_section_heading(doc, 3, "Scope and Limitations")
    doc.add_heading("3.1 Scope", level=2)
    add_table(
        doc,
        ["Area", "Implemented scope"],
        [
            ["Application", "Flutter application for Android with a Flutter Web development/testing target, responsive custom UI, GoRouter navigation, immersive Android presentation, and branded launcher assets."],
            ["City generation", "Seeded city setup; 4–6 problems drawn from 12 issue families; 11 tracked city indicators; starting pressure, concern priorities, and trend/severity variation."],
            ["Candidate generation", "Five unique candidates selected from a seven-portrait pool; compatible fictional names; careers, parties, slogans, hidden tradeoffs, histories, relationships, controversies, promises, and information footprints."],
            ["Candidate investigation", "Candidate Roster; Bio/Civic Archive; four promises; Then-vs-Now context; Campaign Trail; three-topic Debate; nine evidence records per candidate; dossiers; source readers; bookmarks, pins, notes, costs/unlocks, Receipts Board, and comparison."],
            ["Information ecosystem", "Issue-rooted Chronicle and resident voices; candidate and term CityBook posts; three editorial framings of shared events; full News articles; deterministic R-8 TV reports."],
            ["Election", "Game-styled ballot, candidate/issue/confidence reflection, confirmation, and animated winner reveal."],
            ["Living City", "One 365-day term with Day/Week/Month/Year controls, 11 indicators, issue state, promise-derived projects, world conditions, events, causal chains, map markers, 27 semantic anchors, 19 persistent landmarks, history, and final report."],
            ["Local persistence", "Five Sembast save slots with schema-8 migration; generated election snapshot, investigation state, vote, full year history, current day/speed/pause state, inspected events, and settings persistence."],
            ["Community API", "Optional Anonymous Auth identity, 3–24 character display name, Public/Private city snapshot publishing, Explore, My Cities, player search, privacy changes, deletion, separate caches, and explicit REST CRUD."],
            ["Language and audio", "English/Filipino localization; reduced motion; high contrast; SFX/dialogue/music toggles and volumes; route-aware music contexts; critical-event interruption settings."],
        ],
        widths=[1.45, 5.05],
    )

    doc.add_heading("3.2 Limitations", level=2)
    add_bullets(
        doc,
        [
            "The content is entirely fictional and cannot predict real elections, score real politicians, or recommend how a person should vote.",
            "The current loop covers one mayoral election and one 365-day term. Vice Mayor, multiple offices, connected future elections, and long-term administration history are not yet implemented.",
            "The simulation is aggregate and systemic. It does not run individual citizen, vehicle, building, household, or road-agent artificial intelligence.",
            "The player observes the elected administration; the current term does not let the player act as mayor or directly allocate the city budget.",
            "Runtime AI does not determine simulation truth, candidate integrity, evidence, or outcomes. Current variation is produced by deterministic rules and authored content libraries.",
            "There is no real-time multiplayer, friends/follow system, messaging, social comments between real users, rankings, or shared cloud save.",
            "Supabase sharing is asynchronous and optional. Clearing app data or changing devices can lose access to an anonymous identity because no email/password recovery system is implemented.",
            "The deployed Supabase project must still be checked manually for the V2 SQL, Anonymous Sign-Ins, RLS state, and live request behavior; repository tests verify code and intended policy definitions, not dashboard state.",
            "Music tracks and the reporter dialogue sound require verified redistribution provenance or replacement before a public/commercial release. The application remains usable without optional music.",
            "The Android classroom APK currently uses debug signing. A private upload/release keystore is required before Play Store or public distribution.",
            "Some media slots reuse existing art or generic safe fallbacks; the procedural content and visual libraries can be expanded without changing current saved worlds.",
        ],
    )

    doc.add_page_break()
    add_section_heading(doc, 4, "System Architecture")
    add_figure(
        doc,
        ASSET_DIR / "system_architecture.png",
        "Figure 1. Current Resibo, Please system architecture.",
        width=6.55,
    )
    add_paragraph(
        doc,
        "The application follows a local-first layered design. Flutter screens and GoRouter present the game and collect user actions. GameController coordinates the active save and delegates truth generation to deterministic domain engines. Typed models represent the city, candidate, evidence, election, events, and term state. Repositories and codecs persist the resulting state instead of asking widgets to regenerate it."
    )
    add_table(
        doc,
        ["Layer", "Current responsibility", "Representative implementation"],
        [
            ["Presentation", "Responsive game screens, navigation, localization, visual states, accessibility, and input", "lib/features, lib/app/router.dart, lib/app/theme"],
            ["Application state", "Active save selection, investigation mutations, election handoff, simulation visibility, and save scheduling", "lib/core/state/game_controller.dart"],
            ["Domain", "Seeded city/candidate generation, evidence relationships, information ecosystem, election outcome, and Living City rules", "lib/domain/models, lib/domain/simulation"],
            ["Data", "Authored libraries, scenario configuration, persistence codecs, gallery service/repository/cache", "lib/data"],
            ["Local storage", "Canonical city saves, settings, and read-only gallery cache", "Sembast + SharedPreferences"],
            ["Optional online", "Anonymous identity and sanitized community records", "Supabase Auth + PostgREST + PostgreSQL/RLS"],
        ],
        widths=[1.05, 3.25, 2.2],
    )
    add_callout(
        doc,
        "Architecture principle",
        "Algorithm creates truth; presentation displays that truth. Chronicle, News, TV, CityBook, map markers, and the Term Report consume shared saved event or candidate records rather than operating as independent simulation engines.",
        color=TEAL,
    )

    doc.add_page_break()
    add_section_heading(doc, 5, "System Flowchart")
    add_figure(
        doc,
        ASSET_DIR / "system_flowchart.png",
        "Figure 2. Current player and system flow from launch to optional city sharing.",
        width=5.0,
    )
    add_paragraph(
        doc,
        "A player may leave and continue from the current save stage. Before voting, the City, Dossiers, and Chronicle navigation keeps city conditions separate from candidate claims while still allowing cross-links into evidence and media. After voting, the current simulation day determines which events and related Chronicle, News, TV, and CityBook records are visible."
    )

    doc.add_page_break()
    add_section_heading(doc, 6, "Database Design / ERD")
    doc.add_heading("6.1 Online database", level=2)
    add_paragraph(
        doc,
        "The optional community feature uses the Supabase REST API (PostgREST) and Supabase Anonymous Auth. The current V2 SQL defines two public tables linked to Supabase-managed auth.users. The anonymous session receives a stable auth.uid() on the current installation, while the player supplies only a display name."
    )
    add_figure(
        doc,
        ASSET_DIR / "supabase_erd.png",
        "Figure 3. Supabase V2 entity-relationship diagram.",
        width=6.55,
    )
    doc.add_heading("6.2 Current table schema", level=2)
    add_table(
        doc,
        ["Table", "Field", "Type / rule"],
        [
            ["player_profiles", "user_id", "uuid; primary key; defaults to auth.uid(); foreign key to auth.users(id); cascade delete"],
            ["player_profiles", "display_name", "text; required; trimmed; 3–24 characters; supported letters/numbers/underscore/spaces"],
            ["player_profiles", "display_name_normalized", "text; required; unique index; must equal lower(display_name)"],
            ["player_profiles", "created_at / updated_at", "timestamptz; server defaults; update trigger maintains updated_at"],
            ["city_snapshots", "id", "uuid; primary key; generated by gen_random_uuid()"],
            ["city_snapshots", "owner_id", "uuid; required; defaults to auth.uid(); foreign key to player_profiles(user_id); cascade delete"],
            ["city_snapshots", "is_public", "boolean; required; default true"],
            ["city_snapshots", "city_name / term_number / starting_pressure", "bounded text/integer values describing the completed fictional city"],
            ["city_snapshots", "top_issue / selected_candidate_name / candidate_party", "bounded text presentation fields"],
            ["city_snapshots", "final_city_condition / outcome_summary", "enumerated condition and summary limited to 600 characters"],
            ["city_snapshots", "major_issues / major_events", "jsonb arrays; strings only; maximum 8 entries and 2048 bytes each"],
            ["city_snapshots", "created_at", "timestamptz; required; default now()"],
        ],
        widths=[1.25, 2.0, 3.25],
    )

    doc.add_heading("6.3 API operations", level=2)
    add_table(
        doc,
        ["Method", "Resource", "Implemented purpose"],
        [
            ["GET", "/rest/v1/player_profiles", "Restore a profile or perform a server-side case-insensitive partial display-name search."],
            ["POST", "/rest/v1/player_profiles", "Create the current anonymous user’s validated display-name profile."],
            ["GET", "/rest/v1/city_snapshots", "Load public Explore rows, owned rows, or one searched player’s public rows using query filters."],
            ["POST", "/rest/v1/city_snapshots", "Publish one completed sanitized city as Public or Private; id and owner_id are not supplied by the client."],
            ["PATCH", "/rest/v1/city_snapshots?id=eq.<id>", "Change only is_public for a snapshot owned by the current authenticated session."],
            ["DELETE", "/rest/v1/city_snapshots?id=eq.<id>", "Delete an owned uploaded snapshot without deleting the local city save."],
        ],
        widths=[0.75, 2.45, 3.3],
    )

    doc.add_heading("6.4 Data minimization and access control", level=2)
    add_bullets(
        doc,
        [
            "Local gameplay saves are not stored in Supabase. Sembast remains the canonical storage for world seed, candidates, evidence, investigation state, vote, events, and term history.",
            "PublicCitySnapshot cannot serialize the world seed, hidden candidate capabilities, complete save document, bookmarks, private notes, or investigation history.",
            "V2 SQL enables and forces row-level security. Authenticated sessions may read public rows or their own rows; insert/update/delete checks compare owner_id with auth.uid().",
            "Column grants restrict city-snapshot updates to is_public, while the unauthenticated database role receives no table privileges.",
            "The SQL file describes intended policy. The actual deployed dashboard must be manually verified before the demonstration.",
        ],
    )

    doc.add_page_break()
    add_section_heading(doc, 7, "Data and Component Architecture")
    add_figure(
        doc,
        ASSET_DIR / "component_data_flow.png",
        "Figure 4. Seeded data flow from city creation to shared information and term outcomes.",
        width=6.55,
    )
    add_paragraph(
        doc,
        "SeedDerivation creates isolated subsystem paths, so a random draw added to CityBook does not consume the path used for candidate identity, city issues, another candidate, or the Living City. A normal generated election is serialized after creation and restored from that snapshot instead of being regenerated during navigation. Candidate generation is version 1, the current Living City engine is generation version 4, and the current local save schema is version 8."
    )
    add_table(
        doc,
        ["Data source", "Generated or stored output", "Primary consumers"],
        [
            ["World seed + city configuration", "4–6 CityProblem records and 11 starting indicators", "City Brief, Chronicle, candidate issue fit, term engine"],
            ["Candidate truth profile", "Identity, capabilities, history, promises, relationships, controversies, evidence graph", "Biography, Promises, Debate, CityBook, News, Records, election and term engine"],
            ["Player investigation state", "Viewed/unlocked evidence, costs, bookmarks, pins, suspicious flags, private notes", "Dossier, evidence reader, Receipts Board, save restore"],
            ["Confirmed election", "Selected candidate, top issue, confidence, deterministic term seed", "Winner reveal, TermEngine, LivingCityEngine"],
            ["Saved CityEvent", "Date, importance, effects, location, related records, parent chain", "Map, Event History, Chronicle, three News framings, TV, CityBook, Term Report"],
            ["Completed TermResult", "Indicator movement, projects, event history, explanation, final condition", "Term Report and sanitized gallery snapshot builder"],
        ],
        widths=[1.45, 2.85, 2.2],
    )

    doc.add_page_break()
    add_section_heading(doc, 8, "Project Progress Report")
    doc.add_heading("8.1 Current overall status", level=2)
    add_paragraph(
        doc,
        "Resibo, Please is a functional midterm MVP on branch v1.0. The full local loop—from city creation and candidate investigation through election, a deterministic 365-day administration, and the Term Report—is implemented. The optional Public City Gallery is implemented in code with profile setup, Public/Private publishing, Explore, My Cities, search, privacy changes, deletion, and cached read states. Current work is primarily refinement, deployment verification, content expansion, and public-release preparation rather than construction of a missing core loop."
    )
    add_callout(
        doc,
        "Fresh verification — September 2, 2026",
        "Static analysis completed with no issues. The complete Flutter test suite passed all 155 tests. Debug and release APK artifacts are present in build/app/outputs/flutter-apk; the classroom release still uses debug signing.",
        color=GREEN,
    )

    doc.add_heading("8.2 Progress by area", level=2)
    add_table(
        doc,
        ["Area", "Status", "Current accomplishment"],
        [
            ["UI/UX", "Implemented; refinement", "Custom political-cartoon/mobile-game visual system across major screens; responsive Main Menu, safe scrolling, immersive Android layout, and non-default game controls."],
            ["Navigation", "Complete for MVP", "GoRouter connects Main Menu, gallery, City Brief, Chronicle, candidate systems, News, TV, election, simulation, and report routes."],
            ["City setup/generation", "Complete for MVP", "Seven setup choices, deterministic world seed, 4–6 active issues, 12 issue-family library, issue metrics and indicators."],
            ["Candidate generation", "Complete for MVP", "Five-of-seven portraits, unique stable identities, hidden tradeoffs, structured history, promises, CityBook, Debate, News, and 45 evidence records per election."],
            ["Candidate investigation", "Complete for MVP", "Roster, Candidate Hub, Civic Archive, Promises, Campaign Trail, Debate, Dossier, source readers, evidence costs/unlocks, Receipts Board, notes, and comparison."],
            ["Information ecosystem", "Complete for MVP", "Chronicle, resident voices, CityBook, one-event/three-outlet News, full articles, and R-8 reports consume shared records."],
            ["Election", "Complete", "Ballot reflection, confirmation, animated winner reveal, and transfer of selected candidate truth into the term simulation."],
            ["Living City simulation", "Complete for one term; refinement", "365 days, Day/Week/Month/Year controls, event templates, projects, world conditions, causal chains, interactive map, landmarks, markers, history, and visibility gating."],
            ["Local persistence", "Complete for MVP", "Five Sembast slots, schema-8 saves/migrations, generated candidate and term snapshots, settings, and recovery-safe repository behavior."],
            ["Settings/localization", "Complete for current UI", "English/Filipino; reduced motion; high contrast; dialogue/SFX/music toggles and volumes; four critical-event interruption modes."],
            ["Audio", "Functional; asset review", "Central cached SFX service, route-aware music contexts, lifecycle handling and TV ducking. Music and reporter-blip provenance still require confirmation/replacement."],
            ["Supabase REST API", "Implemented; live verification required", "Anonymous ownership, profile setup, sanitized public/private snapshots, GET/POST/PATCH/DELETE, validation, safe errors, and separate caches."],
            ["Android build", "Classroom APK ready; production work remains", "Branded adaptive launcher icon, app label, immersive fullscreen, release APK artifact, and responsive phone tests. Production signing/device regression remain."],
            ["Automated testing", "Passing", "155 tests cover generation, integrity, persistence/migration, information consistency, API/SQL intent, settings, responsive UI, and the ballot-to-report loop."],
            ["Documentation", "Completed by this deliverable", "One consolidated midterm document with architecture, flowchart, ERD, component flow, scope, objectives, and evidence-based progress report."],
        ],
        widths=[1.4, 1.25, 3.85],
    )

    doc.add_heading("8.3 Major completed features", level=2)
    add_bullets(
        doc,
        [
            "Responsive illustrated Main Menu, How to Play, five city slots, Continue/Start behavior, Settings, and local persistence.",
            "Procedural City Brief and Chronicle grounded in selected city issues and indicators.",
            "Five-candidate procedural election with truth-first generation, consistent cross-system references, and no exposed best-candidate score.",
            "Candidate Hub, Civic Archive biographies and citations, promises, Campaign Trail, Debate, records/evidence, CityBook, News, TV, and Receipts Board.",
            "Game-styled election, winner reveal, complete Living City term, interactive overhead map, project/event systems, and final Term Report.",
            "Versioned Sembast save/load and migration, SharedPreferences settings, English/Filipino localization, accessibility options, and centralized audio infrastructure.",
            "Optional Supabase gallery with privacy selection, Explore/My Cities, player search, owner-only privacy/delete actions, and sanitized JSON models.",
            "Android launcher branding, immersive system UI behavior, current release artifact, and automated responsive-screen checks.",
        ],
    )

    doc.add_heading("8.4 Features currently being refined", level=2)
    add_bullets(
        doc,
        [
            "Physical-device visual regression, touch ergonomics, audio loudness, lifecycle behavior, and performance on lower-end Android phones.",
            "Live Supabase verification against the intended V2 tables, Anonymous Sign-Ins, RLS policies, ownership, and all four city-snapshot request methods.",
            "Additional procedural text and visual assets for candidate media, issue files, and term events while preserving stable saved-version behavior.",
            "Fine layout consistency across less frequently visited modals, media readers, and small/large screen sizes.",
            "Audio provenance and replacement of any classroom-only music or reporter sound before public distribution.",
        ],
    )

    doc.add_heading("8.5 Remaining work after the midterm", level=2)
    add_bullets(
        doc,
        [
            "Connect completed terms into multiple election cycles with inherited city state, election history, and end-of-run reflection.",
            "Expand long-term candidate/administration history, laws, promise tracking, and content anti-repetition across elections.",
            "Design moderation and richer comparisons before expanding the public gallery’s social surface.",
            "Add recoverable accounts or cloud saves only if the project later requires cross-device identity; the current local-first model should remain authoritative.",
            "Configure private production signing, finalize licensed audio, optimize release size, and complete store/public-release requirements.",
            "Runtime AI presentation remains deferred. If added later, it must paraphrase structured truth only, have a deterministic fallback, and never decide outcomes or fabricate evidence.",
        ],
    )

    doc.add_heading("8.6 Testing and stability status", level=2)
    add_table(
        doc,
        ["Verification area", "Current evidence", "Boundary"],
        [
            ["Static analysis", "No issues found on the current v1.0 branch on September 2, 2026.", "Does not replace runtime/device testing."],
            ["Automated tests", "155 tests passed, including domain, data, persistence, API/SQL intent, settings, and widget flows.", "Mocked/local tests cannot prove the deployed Supabase dashboard state."],
            ["Gameplay loop", "Widget tests cover the ballot, winner reveal, term events, report, save restore, and targeted navigation flows.", "A full physical-device demonstration path should still be repeated before submission."],
            ["Android artifact", "Debug and release APK files are present; native fullscreen and launcher resources compile.", "Release uses debug signing and is suitable only for classroom/demo distribution."],
            ["Performance design", "One raster map, bounded daily logic, capped marker widgets, cached audio, and saved event records reduce runtime work.", "Low-end phone behavior requires manual confirmation."],
        ],
        widths=[1.25, 3.25, 2.0],
    )

    doc.add_heading("8.7 Repository documentation inconsistencies found", level=2)
    add_paragraph(
        doc,
        "Actual source code was prioritized where older reports disagreed. The current LivingCityEngine generationVersion is 4 even though an older architecture paragraph names version 2. GameSaveData.currentSchemaVersion is 8 even though an earlier candidate-system paragraph still mentions schema 7. The August release-readiness report says no Android SDK/APK was available, but the current workspace contains successfully built debug and release APK artifacts and a branded launcher icon. These are documentation-age differences; the current implementation and this report use the source and verified build state."
    )

    doc.add_page_break()
    add_section_heading(doc, 9, "Summary")
    add_paragraph(
        doc,
        "Resibo, Please currently demonstrates its intended educational core: the player must understand a fictional city, investigate five imperfect candidates through several information formats, make an independent electoral decision, and review how that decision interacts with institutions, budget pressure, public trust, city problems, projects, and external conditions over one simulated year."
    )
    add_paragraph(
        doc,
        "The strongest technical characteristic is consistency. Versioned seeded generation creates structured truth, the information ecosystem presents that truth from different perspectives, local persistence prevents rerolls, and the Living City reuses saved event IDs across the map, Chronicle, News, television, CityBook, and final report. The optional Supabase feature is deliberately smaller than the local save and shares only a sanitized outcome summary."
    )
    add_paragraph(
        doc,
        "For the midterm, the system is a complete one-election MVP with passing automated checks and an Android classroom artifact. The next priorities are physical-device and live-backend verification, media/audio licensing cleanup, and controlled expansion toward multiple connected elections rather than broadening the scope before the current loop is stable."
    )

    doc.add_page_break()
    doc.add_heading("Appendix A — Selected System Screenshots", level=1)
    add_paragraph(
        doc,
        "The following tracked screenshots are actual project captures and remain representative of the current Main Menu and pre-election City Brief. Older repository captures of the authored three-candidate ballot, manual four-phase term, and V1 gallery were intentionally excluded because those interfaces have been superseded."
    )
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.rows[0].cells
    for cell, image_path, label in [
        (left, ROOT / "docs" / "week1" / "actual_screens" / "01_main_menu.png", "Figure A1. Main Menu of Resibo, Please."),
        (right, ROOT / "docs" / "week1" / "actual_screens" / "02_city_brief.png", "Figure A2. Pre-election City Brief."),
    ]:
        set_cell_margins(cell, top=80, bottom=80, start=80, end=80)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(image_path), width=Inches(2.75))
        caption = cell.add_paragraph(label)
        caption.style = doc.styles["Figure Caption"]
    add_paragraph(
        doc,
        "Screenshot source: docs/week1/actual_screens in the current repository. No mock or fabricated application screen is presented in this appendix.",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    ).runs[0].italic = True

    return doc


def main() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    generate_architecture_diagram(ASSET_DIR / "system_architecture.png")
    generate_flow_diagram(ASSET_DIR / "system_flowchart.png")
    generate_erd_diagram(ASSET_DIR / "supabase_erd.png")
    generate_component_diagram(ASSET_DIR / "component_data_flow.png")
    document = build_document()
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")
    for image_path in sorted(ASSET_DIR.glob("*.png")):
        print(f"Created {image_path}")


if __name__ == "__main__":
    main()
